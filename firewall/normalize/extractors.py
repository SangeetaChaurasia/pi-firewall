"""Turn every supported input type into a list of Segments.

Key idea: separate what a HUMAN would see (visible segments) from what only a MACHINE
would read (hidden CSS text, HTML comments, metadata, alt text, tiny/white PDF text...).
Hidden segments are scanned with extra suspicion and never forwarded by default.
"""
from __future__ import annotations

import email
import io
import json
import re
from email import policy
from pathlib import Path

from ..models import Segment, SourceType

MIN_HIDDEN = 8


def _as_text(data: str | bytes) -> str:
    if isinstance(data, str):
        return data
    for enc in ("utf-8-sig", "utf-16"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _as_bytes(data: str | bytes) -> bytes:
    return data.encode("utf-8", "replace") if isinstance(data, str) else data


_EXT_MAP = {
    ".pdf": SourceType.PDF, ".docx": SourceType.DOCX,
    ".html": SourceType.HTML, ".htm": SourceType.HTML,
    ".md": SourceType.MARKDOWN, ".markdown": SourceType.MARKDOWN,
    ".eml": SourceType.EMAIL, ".json": SourceType.API_RESPONSE,
    ".png": SourceType.IMAGE, ".jpg": SourceType.IMAGE, ".jpeg": SourceType.IMAGE,
    ".bmp": SourceType.IMAGE, ".tif": SourceType.IMAGE, ".tiff": SourceType.IMAGE, ".webp": SourceType.IMAGE,
    ".py": SourceType.SOURCE_CODE, ".js": SourceType.SOURCE_CODE, ".ts": SourceType.SOURCE_CODE,
    ".java": SourceType.SOURCE_CODE, ".c": SourceType.SOURCE_CODE, ".cpp": SourceType.SOURCE_CODE,
    ".go": SourceType.SOURCE_CODE, ".rs": SourceType.SOURCE_CODE, ".rb": SourceType.SOURCE_CODE,
    ".php": SourceType.SOURCE_CODE, ".sh": SourceType.SOURCE_CODE, ".sql": SourceType.SOURCE_CODE,
}


def guess_source(filename: str) -> SourceType:
    return _EXT_MAP.get(Path(filename).suffix.lower(), SourceType.USER_MESSAGE)


# --------------------------------------------------------------------------- text
def _text(data, filename=None):
    return [Segment(_as_text(data), "text")], []


# --------------------------------------------------------------------------- HTML
_HIDDEN_STYLE = re.compile(
    r"display\s*:\s*none|visibility\s*:\s*hidden|font-size\s*:\s*0(?:px|pt|em|rem|%)?\s*(?:;|$)"
    r"|opacity\s*:\s*0(?:\.0+)?\s*(?:;|$)|(?:left|top|text-indent)\s*:\s*-\d{3,}"
    r"|height\s*:\s*0(?:px)?\s*;[^\"]*overflow\s*:\s*hidden",
    re.I,
)
_WHITE_TEXT = re.compile(r"(?<![-\w])color\s*:\s*(?:#f{3}(?:f{3})?\b|white\b|rgb\(\s*255\s*,\s*255\s*,\s*255\s*\))", re.I)


def _html_segments(html: str, prefix: str = "html") -> list[Segment]:
    from bs4 import BeautifulSoup, Comment

    try:
        soup = BeautifulSoup(html, "lxml")
    except Exception:  # pragma: no cover
        soup = BeautifulSoup(html, "html.parser")

    hidden: list[Segment] = []

    for c in soup.find_all(string=lambda t: isinstance(t, Comment)):
        t = str(c).strip()
        if len(t) >= MIN_HIDDEN:
            hidden.append(Segment(t, f"{prefix}_comment", hidden=True))
        c.extract()

    for meta in soup.find_all("meta"):
        content = (meta.get("content") or "").strip()
        name = meta.get("name") or meta.get("property") or "meta"
        if len(content) >= MIN_HIDDEN:
            hidden.append(Segment(content, f"{prefix}_meta:{name}", hidden=True))

    for el in soup.find_all(True):
        if getattr(el, "attrs", None) is None:
            continue
        for attr in ("alt", "title", "aria-label", "placeholder"):
            v = el.get(attr)
            if isinstance(v, str) and len(v.strip()) >= MIN_HIDDEN:
                hidden.append(Segment(v.strip(), f"{prefix}_attr:{attr}", hidden=True))
        if el.name == "img":
            src = el.get("src") or ""
            if src.startswith(("http://", "https://")) and "?" in src:
                hidden.append(Segment(f"![{el.get('alt', '')}]({src})", f"{prefix}_img_src", hidden=True))
        if el.name == "input" and (el.get("type") or "").lower() == "hidden":
            v = (el.get("value") or "").strip()
            if len(v) >= MIN_HIDDEN:
                hidden.append(Segment(v, f"{prefix}_hidden_input", hidden=True))

    to_remove = []
    for el in soup.find_all(True):
        if getattr(el, "attrs", None) is None:
            continue
        style = el.get("style") or ""
        if el.has_attr("hidden") or _HIDDEN_STYLE.search(style) or _WHITE_TEXT.search(style):
            to_remove.append(el)
    for el in to_remove:
        try:
            if getattr(el, "decomposed", False) or el.attrs is None:
                continue
            t = el.get_text(" ", strip=True)
            if len(t) >= MIN_HIDDEN:
                hidden.append(Segment(t, f"{prefix}_hidden_element", hidden=True))
            el.decompose()
        except Exception:  # pragma: no cover
            continue

    for tag in soup.find_all(["script", "style", "template"]):
        tag.decompose()

    visible_text = re.sub(r"\n{3,}", "\n\n", soup.get_text("\n", strip=True))
    return [Segment(visible_text, f"{prefix}_visible")] + hidden


def _html(data, filename=None):
    return _html_segments(_as_text(data)), []


# --------------------------------------------------------------------------- Markdown
def _markdown(data, filename=None):
    text = _as_text(data)
    segs: list[Segment] = []
    for m in re.finditer(r"<!--(.*?)-->", text, re.S):
        t = m.group(1).strip()
        if len(t) >= MIN_HIDDEN:
            segs.append(Segment(t, "md_html_comment", hidden=True))
    for m in re.finditer(r"\]\(\s*[^)\s]+\s+\"([^\"]{8,})\"\s*\)", text):
        segs.append(Segment(m.group(1), "md_link_title", hidden=True))
    visible = re.sub(r"<!--.*?-->", "", text, flags=re.S)
    return [Segment(visible, "md_visible")] + segs, []


# --------------------------------------------------------------------------- JSON / API
def _json(data, filename=None):
    text = _as_text(data)
    try:
        obj = json.loads(text)
    except json.JSONDecodeError:
        return [Segment(text, "text")], ["api_response_not_valid_json: scanned as plain text"]

    lines: list[str] = []

    def walk(o, path):
        if isinstance(o, dict):
            for k, v in o.items():
                walk(v, f"{path}.{k}" if path else str(k))
        elif isinstance(o, list):
            for i, v in enumerate(o):
                walk(v, f"{path}[{i}]")
        else:
            val = str(o).replace("\n", " ")
            lines.append(f"{path}: {val}")

    walk(obj, "")
    return [Segment("\n".join(lines), "json_flat")], []


# --------------------------------------------------------------------------- Source code
_CODE_PATTERNS = [
    ("code_comment", re.compile(r'"""(.*?)"""', re.S)),
    ("code_comment", re.compile(r"'''(.*?)'''", re.S)),
    ("code_comment", re.compile(r"/\*(.*?)\*/", re.S)),
    ("code_comment", re.compile(r"(?<![:\w\"'])#(?!!)([^\n]*)")),
    ("code_comment", re.compile(r"(?<![:\w\"'])//([^\n]*)")),
    ("code_comment", re.compile(r"<!--(.*?)-->", re.S)),
    ("code_string", re.compile(r'"([^"\\\n]{20,})"')),
    ("code_string", re.compile(r"'([^'\\\n]{20,})'")),
]


def _code(data, filename=None):
    text = _as_text(data)
    segs = [Segment(text, "code", scan=False)]
    seen = set()
    for origin, rx in _CODE_PATTERNS:
        for m in rx.finditer(text):
            if m.span() in seen or len(segs) > 500:
                continue
            seen.add(m.span())
            body = m.group(1).strip()
            if body:
                segs.append(Segment(body, origin, root=0, root_span=m.span(), output=False))
    return segs, []


# --------------------------------------------------------------------------- OCR helpers
def _ocr_image(img):
    """Returns (visible_text, extra_low_contrast_text, warning)."""
    try:
        import pytesseract
        from PIL import ImageOps
    except ImportError:
        return "", "", "ocr_unavailable: pip install pytesseract pillow"
    try:
        gray = img.convert("L")
        base = pytesseract.image_to_string(gray)
        boosted = pytesseract.image_to_string(ImageOps.equalize(ImageOps.autocontrast(gray, cutoff=1)))
    except Exception as e:  # Tesseract binary missing, corrupt image, ...
        return "", "", f"ocr_failed: {type(e).__name__}: {str(e)[:120]}"
    norm = lambda s: re.sub(r"\W+", " ", s).strip().lower()
    known = {norm(l) for l in base.splitlines()}
    extra = [
        l.strip() for l in boosted.splitlines()
        if len(l.strip()) >= MIN_HIDDEN and len(l.split()) >= 3 and norm(l) not in known
        and sum(c.isalpha() for c in l) / max(len(l), 1) > 0.6
    ]
    return base.strip(), "\n".join(extra), ""


def _image(data, filename=None):
    try:
        from PIL import Image
    except ImportError:
        return [], ["image_support_unavailable: pip install pillow pytesseract"]
    try:
        img = Image.open(io.BytesIO(_as_bytes(data)))
        img.load()
    except Exception as e:
        return [], [f"image_unreadable: {type(e).__name__}"]
    visible, extra, warn = _ocr_image(img)
    segs = []
    if visible:
        segs.append(Segment(visible, "ocr_text"))
    if extra:
        segs.append(Segment(extra, "ocr_low_contrast_text", hidden=True))
    return segs, ([warn] if warn else [])


# --------------------------------------------------------------------------- PDF
def _pdf(data, filename=None):
    try:
        try:
            import pymupdf as fitz
        except ImportError:
            import fitz
    except ImportError:
        return [], ["pdf_support_unavailable: pip install pymupdf"]

    warnings: list[str] = []
    segs: list[Segment] = []
    hidden: list[Segment] = []
    try:
        doc = fitz.open(stream=_as_bytes(data), filetype="pdf")
    except Exception as e:
        return [], [f"pdf_unreadable: {type(e).__name__}"]

    for k, v in (doc.metadata or {}).items():
        if k in ("title", "subject", "keywords", "author", "creator") and isinstance(v, str) and len(v.strip()) >= MIN_HIDDEN:
            hidden.append(Segment(v.strip(), f"pdf_metadata:{k}", hidden=True))

    for pno, page in enumerate(doc, start=1):
        vis_lines, hid_lines = [], []
        for block in page.get_text("dict").get("blocks", []):
            if block.get("type") != 0:
                continue
            for line in block.get("lines", []):
                vis, hid = [], []
                for sp in line.get("spans", []):
                    t = sp.get("text", "")
                    if not t.strip():
                        vis.append(t)
                        continue
                    off_page = not page.rect.intersects(fitz.Rect(sp["bbox"]))
                    if sp.get("size", 10) < 2.0 or (sp.get("color", 0) & 0xFFFFFF) == 0xFFFFFF \
                            or sp.get("alpha", 255) == 0 or off_page:
                        hid.append(t)
                    else:
                        vis.append(t)
                if "".join(vis).strip():
                    vis_lines.append("".join(vis))
                if "".join(hid).strip():
                    hid_lines.append("".join(hid))
        try:
            for a in page.annots() or []:
                c = (a.info or {}).get("content", "")
                if len(c.strip()) >= MIN_HIDDEN:
                    hidden.append(Segment(c.strip(), f"pdf_annotation:p{pno}", hidden=True))
        except Exception:  # pragma: no cover
            pass

        vis_text = "\n".join(vis_lines)
        if len(vis_text.strip()) < 20 and page.get_images():
            try:
                from PIL import Image
                pix = page.get_pixmap(dpi=150, alpha=False)
                img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
                ocr_vis, ocr_extra, warn = _ocr_image(img)
                if warn:
                    warnings.append(f"page {pno}: {warn}")
                if ocr_vis:
                    vis_text = (vis_text + "\n" + ocr_vis).strip()
                if ocr_extra:
                    hidden.append(Segment(ocr_extra, f"pdf_ocr_low_contrast:p{pno}", hidden=True))
            except Exception as e:
                warnings.append(f"page {pno}: scanned page OCR failed ({type(e).__name__})")
        if vis_text.strip():
            segs.append(Segment(vis_text, f"pdf_page:{pno}"))
        if hid_lines:
            hidden.append(Segment("\n".join(hid_lines), f"pdf_hidden_text:p{pno}", hidden=True))
    return segs + hidden, warnings


# --------------------------------------------------------------------------- DOCX
def _run_hidden(run) -> bool:
    f = run.font
    if f.hidden:
        return True
    try:
        if f.size is not None and f.size.pt < 2:
            return True
        if f.color is not None and f.color.type is not None and str(f.color.rgb).upper() == "FFFFFF":
            return True
    except Exception:  # pragma: no cover
        pass
    return False


def _docx(data, filename=None):
    try:
        import docx
    except ImportError:
        return [], ["docx_support_unavailable: pip install python-docx"]
    try:
        d = docx.Document(io.BytesIO(_as_bytes(data)))
    except Exception as e:
        return [], [f"docx_unreadable: {type(e).__name__}"]

    vis_parts: list[str] = []
    hid_parts: list[str] = []

    def handle(p):
        if not p.runs:
            if p.text.strip():
                vis_parts.append(p.text)
            return
        v, h = [], []
        for r in p.runs:
            (h if (r.text.strip() and _run_hidden(r)) else v).append(r.text)
        if "".join(v).strip():
            vis_parts.append("".join(v))
        if "".join(h).strip():
            hid_parts.append("".join(h))

    for s in d.sections:
        for p in s.header.paragraphs:
            handle(p)
    for p in d.paragraphs:
        handle(p)
    for t in d.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    handle(p)
    for s in d.sections:
        for p in s.footer.paragraphs:
            handle(p)

    segs = [Segment("\n".join(vis_parts), "docx_visible")]
    if hid_parts:
        segs.append(Segment("\n".join(hid_parts), "docx_hidden_text", hidden=True))
    cp = d.core_properties
    for name in ("title", "subject", "keywords", "comments", "author", "category"):
        v = getattr(cp, name, "") or ""
        if len(v.strip()) >= MIN_HIDDEN:
            segs.append(Segment(v.strip(), f"docx_metadata:{name}", hidden=True))
    try:
        for c in d.comments:
            if len(c.text.strip()) >= MIN_HIDDEN:
                segs.append(Segment(c.text.strip(), "docx_comment", hidden=True))
    except Exception:
        pass
    return segs, []


# --------------------------------------------------------------------------- Email
def _merge(dst: list[Segment], sub: list[Segment]) -> None:
    off = len(dst)
    for s in sub:
        if s.root is not None:
            s.root += off
        dst.append(s)


def _email(data, filename=None):
    raw = _as_bytes(data)
    msg = email.message_from_bytes(raw, policy=policy.default)
    if not (msg["From"] or msg["Subject"] or msg["To"]):
        return [Segment(_as_text(data), "text")], ["email_headers_not_found: scanned as plain text"]

    segs: list[Segment] = []
    warnings: list[str] = []
    plain: list[str] = []
    subject = str(msg["Subject"] or "").strip()
    if subject:
        plain.append(f"Subject: {subject}")
    for part in msg.walk():
        if part.is_multipart():
            continue
        ctype = part.get_content_type()
        fname = part.get_filename()
        try:
            if fname:
                payload = part.get_payload(decode=True) or b""
                src = guess_source(fname)
                if src in (SourceType.PDF, SourceType.DOCX, SourceType.IMAGE, SourceType.HTML,
                           SourceType.MARKDOWN, SourceType.API_RESPONSE):
                    sub, w = extract(payload, src, fname)
                    for s in sub:
                        s.origin = f"attachment:{fname}/{s.origin}"
                    _merge(segs, sub)
                    warnings += [f"{fname}: {x}" for x in w]
                continue
            if ctype == "text/plain":
                plain.append(part.get_content())
            elif ctype == "text/html":
                sub = _html_segments(part.get_content(), prefix="email_html")
                _merge(segs, sub)
        except Exception as e:  # pragma: no cover
            warnings.append(f"email part skipped ({type(e).__name__})")
    body = Segment("\n".join(plain), "email_body")
    return [body] + segs, warnings


# --------------------------------------------------------------------------- dispatcher
def extract(data, source: SourceType, filename: str | None = None) -> tuple[list[Segment], list[str]]:
    table = {
        SourceType.USER_MESSAGE: _text, SourceType.OCR_TEXT: _text,
        SourceType.WEB_PAGE: _html, SourceType.HTML: _html,
        SourceType.MARKDOWN: _markdown, SourceType.EMAIL: _email,
        SourceType.API_RESPONSE: _json, SourceType.SOURCE_CODE: _code,
        SourceType.PDF: _pdf, SourceType.DOCX: _docx, SourceType.IMAGE: _image,
    }
    return table[source](data, filename)
